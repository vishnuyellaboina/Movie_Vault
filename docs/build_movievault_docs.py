from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\Movie_Vault")
DOCS_DIR = ROOT / "docs"
ASSETS_DIR = ROOT / "assets"
OUTPUT_DOCX = DOCS_DIR / "MovieVault_Project_Documentation.docx"

ACCENT = RGBColor(27, 144, 255)
NAVY = RGBColor(11, 16, 32)
MID = RGBColor(78, 89, 110)
LIGHT = RGBColor(242, 245, 250)
GRAY_BORDER = "D9E1EA"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color=GRAY_BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def style_run(run, size, bold=False, color=None, font="Aptos"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    if color:
        run.font.color.rgb = color


def add_paragraph(doc, text="", style=None, align=None, space_before=0, space_after=0, line_spacing=1.2):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing
    return p


def add_heading(doc, text, level=1):
    p = add_paragraph(doc, space_before=8 if level == 1 else 4, space_after=4)
    run = p.add_run(text)
    style_run(run, 18 if level == 1 else 13, bold=True, color=NAVY)
    return p


def add_body(doc, text, bullet=False):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.22
    run = p.add_run(text)
    style_run(run, 10.5, color=MID)
    return p


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.15)
    table.columns[1].width = Inches(4.1)
    for left, right in rows:
        cells = table.add_row().cells
        for idx, value in enumerate((left, right)):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            style_run(run, 10, bold=idx == 0, color=NAVY if idx == 0 else MID)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[idx])
            if idx == 0:
                set_cell_shading(cells[idx], "EDF4FB")
    return table


def add_feature_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    headers = ["Area", "Current implementation", "Product value"]
    for cell, value in zip(hdr, headers):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(value)
        style_run(run, 10.5, bold=True, color=NAVY)
        set_cell_shading(cell, "DCEBFA")
        set_cell_border(cell, "BDD5EE")
    widths = [Inches(1.45), Inches(2.55), Inches(2.5)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    for area, impl, value in rows:
        cells = table.add_row().cells
        for idx, content in enumerate((area, impl, value)):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(content)
            style_run(run, 9.5, color=MID if idx else NAVY, bold=idx == 0)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[idx])
    return table


def add_two_col_table(doc, rows, headers=("Directory", "Purpose")):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for cell, value in zip(hdr, headers):
        p = cell.paragraphs[0]
        run = p.add_run(value)
        style_run(run, 10.5, bold=True, color=NAVY)
        set_cell_shading(cell, "DCEBFA")
        set_cell_border(cell, "BDD5EE")
    widths = [Inches(1.75), Inches(4.75)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    for left, right in rows:
        cells = table.add_row().cells
        for idx, content in enumerate((left, right)):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(content)
            style_run(run, 9.5, color=MID if idx else NAVY, bold=idx == 0)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[idx])
    return table


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F5F9FE")
    set_cell_border(cell, "C7DBF1")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    title_run = p.add_run(title + "\n")
    style_run(title_run, 10.5, bold=True, color=NAVY)
    body_run = p.add_run(body)
    style_run(body_run, 9.5, color=MID)
    return table


def add_labeled_item(doc, label, text, bullet=False):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.22
    lead = p.add_run(label + ": ")
    style_run(lead, 10.5, bold=True, color=NAVY)
    body = p.add_run(text)
    style_run(body, 10.5, color=MID)
    return p


def make_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)

    logo = ASSETS_DIR / "movievault-logo.png"
    if logo.exists():
        p = add_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
        p.add_run().add_picture(str(logo), width=Inches(2.65))

    title = add_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    style_run(title.add_run("MovieVault"), 26, bold=True, color=NAVY)
    subtitle = add_paragraph(
        doc,
        "Project documentation and engineering handoff",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=10,
    )
    style_run(subtitle.runs[0], 12, color=ACCENT)

    summary = add_paragraph(
        doc,
        "MovieVault is a production-minded Expo mobile application for discovering movies across multiple industries using TMDb as the primary content source and OMDb for high-value enrichment such as IMDb ratings and box office information.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
        line_spacing=1.3,
    )
    style_run(summary.runs[0], 11, color=MID)

    add_labeled_item(doc, "Project type", "React Native mobile app built with Expo SDK 54")
    add_labeled_item(doc, "Primary platforms", "Android-first Expo Go development flow with EAS-ready build setup")
    add_labeled_item(doc, "Data providers", "TMDb for discovery, credits, providers, and imagery; OMDb for IMDb rating and box office")
    add_labeled_item(doc, "App package", "com.movievault.app")
    add_labeled_item(doc, "Version", "1.0.0")
    add_labeled_item(doc, "Current workspace", r"F:\Movie_Vault")

    add_paragraph(doc, space_after=6)
    add_paragraph(doc, space_after=4)
    add_body(
        doc,
        "Documentation intent: this document captures the current implementation, technical architecture, product decisions, deployment setup, and recommended next steps for MovieVault as it exists in the codebase on April 27, 2026.",
    )

    doc.add_page_break()

    add_heading(doc, "1. Product Overview", 1)
    add_body(
        doc,
        "MovieVault is designed as a search-first movie information experience that brings discovery, industry browsing, cast navigation, and movie detail enrichment into a single mobile app.",
    )
    add_body(doc, "The current product direction emphasizes speed, clean mobile presentation, and real metadata rather than placeholder content.", bullet=True)
    add_body(doc, "The home experience is curated around trending rails and industry-specific sections such as Hollywood, Bollywood, Tollywood, Kollywood, and Mollywood.", bullet=True)
    add_body(doc, "Movie detail pages enrich a title with ratings, runtime, cast, streaming providers, and production companies.", bullet=True)
    add_body(doc, "Actor pages surface profile information and filmography using both cast and crew credits, which is important for directors and writers such as Trivikram Srinivas.", bullet=True)

    add_heading(doc, "2. Core Features", 1)
    add_labeled_item(doc, "Home", "Trending rail, multi-industry rails, themed hero area, and a search entry point that keeps discovery lightweight.")
    add_labeled_item(doc, "Search", "Debounced global search, recent search history, and infinite query pagination for broader result coverage.")
    add_labeled_item(doc, "Industry pages", "Dedicated wood screens with limited initial load and a wider industry search when the user starts typing.")
    add_labeled_item(doc, "Movie detail", "Poster or backdrop hero, runtime, ratings, production, streaming providers, top cast, and box office enrichment.")
    add_labeled_item(doc, "Actor detail", "Personal profile, biography, debut movie, top loaded box office title, and 3-column filmography.")
    add_labeled_item(doc, "Theme support", "Dark and light modes with custom home interactions and screen-level theme treatment.")

    add_heading(doc, "3. Technology Stack", 1)
    add_body(doc, "Framework: Expo SDK 54 with React Native 0.81.5 and React 19.1.0.")
    add_body(doc, "Navigation: Expo Router using the file-based app directory.")
    add_body(doc, "Data and caching: TanStack Query 5 with custom service-layer caching and local persistence patterns.")
    add_body(doc, "UI utilities: Expo Image, Expo Linear Gradient, FlashList, AsyncStorage, Lucide-style iconography.")
    add_body(doc, "Build and release: Expo App Services configuration present in eas.json with development, preview, and production profiles.")

    add_heading(doc, "4. Screen Inventory", 1)
    add_labeled_item(doc, "app/index.tsx", "Home screen with themed hero, animated top element, rails, and tap-through search bar.")
    add_labeled_item(doc, "app/search.tsx", "Global search experience with recent history and paginated results.")
    add_labeled_item(doc, "app/category/[industry].tsx", "Industry-specific catalog page with scoped search.")
    add_labeled_item(doc, "app/movie/[id].tsx", "Movie detail page with providers above cast and a production section.")
    add_labeled_item(doc, "app/actor/[id].tsx", "Actor and crew profile page with filmography.")
    add_labeled_item(doc, "app/watchlist.tsx", "Watchlist route shell for future persistent user flows.")

    add_paragraph(doc, space_after=3)
    add_body(
        doc,
        "Current product note: the watchlist route exists in the navigation shell, but it should still be treated as an expansion path rather than a fully mature synchronized feature.",
    )

    add_heading(doc, "5. Data Model and API Flow", 1)
    add_body(doc, "MovieVault uses TMDb as the authoritative source for discovery, credits, provider availability, imagery, and most detail-page metadata.")
    add_body(doc, "OMDb is called selectively for enrichment, mainly to surface IMDb rating and box office values where available.")
    add_body(doc, "The project normalizes remote responses into stable internal models: MovieSummary, MovieDetail, PersonSummary, and PersonDetail.")
    add_body(doc, "For actor pages, cast and crew movie credits are merged and deduplicated before rendering filmography so creators outside acting roles are still represented correctly.")

    add_labeled_item(doc, "MovieSummary", "Lightweight card model used across home, search, and industry rails.")
    add_labeled_item(doc, "MovieDetail", "Expanded title model with runtime, IMDb rating, providers, production companies, and curated people rail.")
    add_labeled_item(doc, "PersonSummary", "Compact cast or crew representation used in movie detail.")
    add_labeled_item(doc, "PersonDetail", "Profile, biography, debut movie, top loaded box office title, and filmography grid.")

    add_heading(doc, "6. Performance and UX Patterns", 1)
    add_body(doc, "The app is intentionally structured to feel fast even when upstream APIs are variable.")
    add_body(doc, "Debounced search reduces API churn and improves perceived responsiveness.", bullet=True)
    add_body(doc, "Parallel detail fetches reduce waterfalls on movie pages.", bullet=True)
    add_body(doc, "Home and category pages avoid loading overly large payloads up front.", bullet=True)
    add_body(doc, "Caching is used in both TanStack Query and service helpers to preserve speed and reduce repeat network work.", bullet=True)
    add_body(doc, "Search history is persisted locally for quick repeat actions.", bullet=True)
    add_body(doc, "Skeleton-first and themed UI decisions reduce the feeling of blank-state latency.", bullet=True)

    add_heading(doc, "7. Project Structure", 1)
    add_labeled_item(doc, "app", "Route-level screens including home, search, category, movie, actor, and watchlist.")
    add_labeled_item(doc, "components", "Reusable cards, lists, filters, search elements, and UI primitives.")
    add_labeled_item(doc, "hooks", "Query and interaction hooks such as useHomeFeed, useMovieSearch, useMovieDetail, and search history.")
    add_labeled_item(doc, "services", "TMDb client, OMDb client, normalization, HTTP helpers, and movie orchestration.")
    add_labeled_item(doc, "types", "Shared TypeScript contracts for movies and people.")
    add_labeled_item(doc, "theme", "Color tokens and theme-aware UI surface definitions.")

    add_heading(doc, "8. Build and Deployment", 1)
    add_body(doc, "Local development is currently optimized for Expo Go on Android. The application is configured to use automatic theme handling and is already wired with Expo Router and EAS metadata.")
    add_body(doc, "Recommended developer flow:", bullet=False)
    add_body(doc, r"1. cd F:\Movie_Vault", bullet=True)
    add_body(doc, "2. npm.cmd install", bullet=True)
    add_body(doc, "3. npm.cmd run start", bullet=True)
    add_body(doc, "4. Open the QR session in Expo Go", bullet=True)
    add_body(doc, "5. For builds, use EAS profiles defined in eas.json", bullet=True)

    add_labeled_item(doc, "Preview build intent", "Internal Android testing and shareable install flow.")
    add_labeled_item(doc, "Production build intent", "Release-ready store artifact pipeline.")
    add_labeled_item(doc, "Current icon and splash", "Configured from assets/movievault-logo.png in app.json.")
    add_labeled_item(doc, "User interface style", "automatic.")

    add_heading(doc, "9. Security and Operational Notes", 1)
    add_body(doc, "The current app.json stores TMDb and OMDb keys in expo.extra, which is acceptable for local prototyping but not ideal for a production-grade public client.")
    add_body(doc, "The architecture notes already point toward the correct next move: place API access behind a Node service with Redis caching.", bullet=True)
    add_body(doc, "That backend layer would improve key protection, cache sharing, rate-limit resilience, and normalization consistency.", bullet=True)
    add_body(doc, "Streaming provider availability is region-sensitive and depends on TMDb watch provider coverage.", bullet=True)

    add_heading(doc, "10. Recommended Next Steps", 1)
    add_body(doc, "Move API traffic behind a backend aggregator with Redis.")
    add_body(doc, "Implement a real persisted watchlist and optional authentication.")
    add_body(doc, "Persist TanStack Query cache for stronger offline behavior.")
    add_body(doc, "Add instrumentation for search latency, cache hit rate, and transition time.")
    add_body(doc, "Create a final build pipeline with tested icons, splash assets, and Play Store-ready metadata.")

    add_paragraph(doc, space_after=4)
    add_body(
        doc,
        "Bottom line: MovieVault already has the skeleton of a serious movie discovery product with normalized data, themed mobile UI, search-first flows, industry-specific browsing, and a codebase arranged for future scale. The biggest remaining step is moving from direct-client API usage to a production backend layer.",
    )

    section = doc.sections[-1]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("MovieVault Project Documentation")
    style_run(run, 9, color=MID)

    doc.save(OUTPUT_DOCX)


if __name__ == "__main__":
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    make_doc()
